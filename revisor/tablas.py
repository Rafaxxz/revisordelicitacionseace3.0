"""Lectura tolerante de CSV / XLSX y conversión de valores."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Optional

from .categorias import normalizar


def leer_filas(ruta: Path | str, contenido: bytes | None = None) -> list[dict[str, str]]:
    """Lee un .csv/.txt/.xlsx y devuelve filas como dict con claves normalizadas."""
    ruta = Path(ruta)
    if contenido is None:
        contenido = ruta.read_bytes()
    ext = ruta.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        filas = _leer_xlsx(contenido)
    else:
        filas = _leer_csv(contenido)
    return filas


def _clave(encabezado: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalizar(str(encabezado or ""))).strip("_")


def _leer_csv(contenido: bytes) -> list[dict[str, str]]:
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:5000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;|\t")
    except csv.Error:
        dialecto = csv.excel
    lector = csv.reader(io.StringIO(texto), dialecto)
    return _a_dicts(lector)


def _leer_xlsx(contenido: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    return _a_dicts(hoja.iter_rows(values_only=True))


def _a_dicts(filas: Iterable[Iterable[object]]) -> list[dict[str, str]]:
    encabezados: list[str] | None = None
    salida = []
    for fila in filas:
        valores = ["" if v is None else v for v in fila]
        if encabezados is None:
            # La primera fila con al menos 3 celdas no vacías es el encabezado
            # (los reportes del Estado suelen traer títulos arriba).
            if sum(1 for v in valores if str(v).strip()) >= 3:
                encabezados = [_clave(v) for v in valores]
            continue
        if not any(str(v).strip() for v in valores):
            continue
        salida.append(
            {k: (v if isinstance(v, (date, datetime)) else str(v).strip())
             for k, v in zip(encabezados, valores) if k}
        )
    return salida


def campo(fila: dict, *candidatos: str) -> str:
    """Primer valor no vacío entre columnas cuyo nombre contenga algún candidato."""
    for cand in candidatos:
        if cand in fila and str(fila[cand]).strip():
            return fila[cand]
    for cand in candidatos:
        for k, v in fila.items():
            # "ganador" no debe capturar la columna "ruc_ganador".
            if "ruc" not in cand and k.startswith("ruc"):
                continue
            if cand in k and str(v).strip():
                return v
    return ""


def a_fecha(valor: object) -> Optional[date]:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor or "").strip()
    if not s:
        return None
    s = s.replace("T", " ").split(" ")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def a_monto(valor: object) -> Optional[float]:
    if isinstance(valor, (int, float)):
        return float(valor)
    s = re.sub(r"[^\d,.\-]", "", str(valor or ""))
    if not s:
        return None
    if "," in s and "." in s:
        # El último separador es el decimal.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        partes = s.split(",")
        s = s.replace(",", ".") if len(partes[-1]) in (1, 2) else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def solo_digitos(valor: object) -> str:
    return re.sub(r"\D", "", str(valor or ""))
